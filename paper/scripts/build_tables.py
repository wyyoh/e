"""Create editable three-rule Word/LaTeX table examples from frozen visual data.
This is a style specimen, not the official competition manuscript template.
No optimization is performed. Never treats unlike resources as a common cost.
"""
from pathlib import Path
import argparse, json, statistics, hashlib
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]

def main(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    dst=out/'editable_tables.docx'
    if dst.exists(): raise FileExistsError(dst)
    data=json.loads((ROOT/'data/visual_data.json').read_text())
    pairs=data['pricing']; bs=sum(x['baseline_median_s'] for x in pairs); ac=sum(x['accelerated_median_s'] for x in pairs)
    labelb=sum(x['baseline_labels'] for x in pairs); labela=sum(x['accelerated_labels'] for x in pairs)
    tables=[]
    tables.append({'title':'表1　固定定价输入上的计算开销对照','headers':['指标','原定价器','加速定价器','比较结果'],
      'rows':[['各输入中位耗时合计 / s',f'{bs:.6f}',f'{ac:.6f}',f'速度比 {bs/ac:.3f}'],
      ['累计生成标签数',f'{labelb:,}',f'{labela:,}',f'减少 {(1-labela/labelb)*100:.2f}%'],
      ['24组最小约化成本','逐组相同','逐组相同','未改变定价最小值'],
      ['耗时较原版更长的输入数','—',str(sum(x['accelerated_median_s']>x['baseline_median_s'] for x in pairs)),'没有删除负向结果']],
      'note':'注：8个固定证书节点×3种机型，共24个定价输入；每输入各运行3次，先求中位数再相加。三次为技术重复，不是三个独立问题实例。速度比不是整个全局求解器的提速倍数。耗时来自已冻结实验记录，本次仅重新制表。',
      'widths':[62,31,31,36]})
    q=data['q3_sampling']; rows=[]
    for key,name in [('uniform','均匀采样'),('feasible','必要条件筛选'),('informed','关键点引导')]:
      vals=[x['sample_best_extra_db'] for x in q if x['method']==key]
      rows.append([name,f'{max(vals):.6f}',f'{statistics.median(vals):.6f}',f'{min(vals):.6f}',f'{sum(v>=.5 for v in vals)}/5'])
    tables.append({'title':'表2　历史Q3同预算采样的静态代理评分','headers':['方法','最好 / dB','中位 / dB','最差 / dB','达0.5 dB'], 'rows':rows,
      'note':'注：这是pre-V6历史筛查实验。5个配对随机种子，每方法每种子192次候选评估；评分假设中继始终在线，不等于连续通信或联合调度保证，更不能作为V6通信裕度。',
      'widths':[42,29,29,29,31]})
    r=data['q4']; rows=[]
    for i,name in enumerate(r['resource_names']):
      rows.append([name,str(r['inventory'][i]),str(r['pooled'][i]),str(r['strict_2'][i]),str(max(0,r['strict_2'][i]-r['inventory'][i])),str(r['strict_3'][i]),str(max(0,r['strict_3'][i]-r['inventory'][i]))])
    tables.append({'title':'表3　严格分区下的分类型资源需求与缺口','headers':['资源类型','库存','共享\n最少','两组\n需求','两组\n缺口','三组\n需求','三组\n缺口'],'rows':rows,
      'note':'注：运输机、中继机单位为架；电池、能源组件单位为组。两组固定为除S006外14区 / S006；三组为S001 / 其余13区 / S006。均冻结同一Q3时间表和实际保障关系。资源不跨组调配、不同类型不可抵扣；本表不合计为采购成本，复制中继情景未混入。',
      'widths':[40,20,20,20,20,20,20]})
    doc=Document(); sec=doc.sections[0]; sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(22);sec.bottom_margin=Mm(22);sec.left_margin=Mm(25);sec.right_margin=Mm(25)
    st=doc.styles['Normal'];st.font.name='Noto Serif CJK SC';st.font.size=Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.paragraph_format.space_after=Pt(5)
    for nm in ['Title','Heading 1','Heading 2']:
      st=doc.styles[nm];st.font.name='Noto Sans CJK SC';st.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    doc.core_properties.title='D题可编辑三线表示例';doc.core_properties.author='';doc.core_properties.last_modified_by=''
    h=sec.header.paragraphs[0];h.text='图表制作样张  |  非比赛官方模板';h.style=doc.styles['Caption']
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    doc.add_heading('可编辑表格样张',0)
    doc.add_paragraph('这些表格从冻结数据自动生成，展示三线表、单位、统计口径和结论边界。无新增求解或实验。正式论文须套用当届官方模板。')
    for ti,t in enumerate(tables):
      if ti==2:
        doc.add_page_break();doc.add_heading('异构资源不合并成单一成本',1)
      p=doc.add_paragraph(t['title']);p.paragraph_format.keep_with_next=True
      for rr in p.runs:rr.bold=True
      table=doc.add_table(rows=1,cols=len(t['headers']));table.autofit=False
      pr=table._tbl.tblPr;bd=OxmlElement('w:tblBorders')
      for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil')
        if edge in ['top','bottom']:x.set(qn('w:sz'),'8')
        bd.append(x)
      pr.append(bd)
      for j,w in enumerate(t['widths']):table.columns[j].width=Mm(w)
      for j,text in enumerate(t['headers']):table.rows[0].cells[j].text=text
      trpr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
      for cell in table.rows[0].cells:
        cb=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');cb.append(b);cell._tc.get_or_add_tcPr().append(cb)
      for row in t['rows']:
        cells=table.add_row().cells
        for c,text in zip(cells,row):c.text=text
      for i,row in enumerate(table.rows):
        rp=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');rp.append(no)
        for j,cell in enumerate(row.cells):
          cell.width=Mm(t['widths'][j])
          for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.RIGHT
            for rr in p.runs:rr.font.size=Pt(9);rr.bold=(i==0)
      p=doc.add_paragraph(t['note']);p.paragraph_format.space_after=Pt(12)
      for rr in p.runs:rr.font.size=Pt(9)
    doc.add_heading('在正式论文中沿用的表格规则',2)
    doc.add_paragraph('表题置表上，图题置图下；单位放表头，含义及边界写表注；数值列统一精度并右对齐。0表示真零，破折号表示不适用，未知或限时须直接写出。表格不使用截图，不依赖背景色标示优劣。正式表格编号应使用自动题注与交叉引用。')
    doc.save(dst)
    def esc(s):return str(s).replace('&',r'\&').replace('%',r'\%').replace('_',r'\_').replace('\n',' ')
    tex=[]
    for i,t in enumerate(tables):
      tex.append(r'\begin{table}[htbp]\centering\small'+'\n'+r'\caption{'+esc(t['title'].split('　',1)[1])+'}\n'+r'\begin{tabular}{l'+'r'*(len(t['headers'])-1)+'}\n'+r'\toprule'+'\n')
      tex.append(' & '.join(map(esc,t['headers']))+r' \\'+'\n'+r'\midrule'+'\n')
      tex += [' & '.join(map(esc,row))+r' \\'+'\n' for row in t['rows']]
      tex.append(r'\bottomrule\end{tabular}'+'\n'+r'\par\smallskip\begin{minipage}{\linewidth}\footnotesize '+esc(t['note'])+r'\end{minipage}\end{table}'+'\n')
    (out/'table_examples.tex').write_text(''.join(tex),encoding='utf-8')
    (out/'table_records.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (out/'TABLE_MANIFEST.json').write_text(json.dumps({'source_sha256':hashlib.sha256((ROOT/'data/visual_data.json').read_bytes()).hexdigest(),'native_word_tables':len(doc.tables),'new_optimization':False,'files':[{ 'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [dst,out/'table_examples.tex',out/'table_records.json']]},indent=2)+'\n')
    print(dst)
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--out',type=Path,default=ROOT/'tables/generated');a=ap.parse_args();main(a.out)
